from __future__ import annotations

import asyncio
import logging
from datetime import datetime

import click
import structlog

from src.config import settings
from src.pipeline import run_search
from src.scrapers import SCRAPERS
from src.scrapers.base import SearchParams


def _setup_logging() -> None:
    logging.basicConfig(level=settings.log_level.upper(), format="%(message)s")
    structlog.configure(
        processors=[
            structlog.processors.TimeStamper(fmt="iso"),
            structlog.processors.add_log_level,
            structlog.dev.ConsoleRenderer(),
        ]
    )


@click.group()
def cli() -> None:
    """housing-scraper — short-term rental aggregator (Excel output)."""
    _setup_logging()


@cli.command("run")
@click.option("--city", required=True)
@click.option("--state", required=True, help="2-letter state code, e.g. NY")
@click.option(
    "--start", "start_date", required=True,
    type=click.DateTime(formats=["%Y-%m-%d"]),
    help="Move-in date YYYY-MM-DD",
)
@click.option(
    "--end", "end_date", required=True,
    type=click.DateTime(formats=["%Y-%m-%d"]),
    help="Move-out date YYYY-MM-DD",
)
@click.option(
    "--output", "-o", default="listings.xlsx", show_default=True,
    help="Path to write the .xlsx file.",
)
@click.option("--max-price", "max_price", type=int, default=None)
@click.option("--bedrooms", "bedrooms", type=int, default=None)
@click.option(
    "--max-results", "max_results", type=int, default=20, show_default=True,
    help="Cap listings yielded per source. Use 0 for no cap.",
)
@click.option(
    "--max-min-stay", "max_min_stay", type=int, default=60, show_default=True,
    help="Skip listings whose minimum stay exceeds this many days. Use 0 for no cap.",
)
@click.option(
    "--exhaustive", is_flag=True, default=False,
    help="Bypass Blueground's 50-result cap by querying every neighborhood "
    "placeId on the city page. Slower but much more coverage.",
)
@click.option(
    "--source", "sources", multiple=True,
    type=click.Choice(list(SCRAPERS.keys())),
    help="Limit to specific source(s); default: blueground",
)
def run_cmd(
    city: str,
    state: str,
    start_date: datetime,
    end_date: datetime,
    output: str,
    max_price: int | None,
    bedrooms: int | None,
    max_results: int,
    max_min_stay: int,
    exhaustive: bool,
    sources: tuple[str, ...],
) -> None:
    """Scrape listings and write them to an Excel workbook."""
    params = SearchParams(
        city=city,
        state=state.upper(),
        start_date=start_date.date(),
        end_date=end_date.date(),
        max_price_monthly_usd=max_price,
        min_bedrooms=bedrooms,
        max_results=max_results or None,
        max_min_stay_days=max_min_stay or None,
        exhaustive=exhaustive,
    )
    src_list = list(sources) if sources else ["blueground"]
    rows = asyncio.run(run_search(params, src_list))

    if not rows:
        click.echo("(no listings found)")
        return

    _write_excel(rows, output)
    click.echo(f"Wrote {len(rows)} listings -> {output}")


def _write_excel(records: list[dict], path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    columns = [
        ("Property", "title", 36),
        ("Beds", "bedrooms", 6),
        ("Baths", "bathrooms", 6),
        ("Sqft", "sqft", 7),
        ("Price ($/mo)", "price_monthly_usd", 13),
        ("Move-in", "available_from", 12),
        ("Min stay (days)", "min_stay_days", 14),
        ("Address", "address", 42),
        ("City", "city", 14),
        ("State", "state", 6),
        ("Lat", "lat", 10),
        ("Lon", "lon", 10),
        ("Source", "source", 16),
        ("Link", "url", 50),
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Listings"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_align = Alignment(horizontal="left", vertical="center")
    thin = Side(border_style="thin", color="E5E7EB")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)

    for col_idx, (label, _, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 22

    records = sorted(
        records,
        key=lambda r: (r.get("price_monthly_usd") is None, r.get("price_monthly_usd") or 0),
    )

    for r_idx, rec in enumerate(records, start=2):
        for c_idx, (_, key, _) in enumerate(columns, start=1):
            val = rec.get(key)
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border
            if key == "url" and val:
                cell.hyperlink = val
                cell.font = Font(color="2563EB", underline="single")
            elif key == "price_monthly_usd" and val is not None:
                cell.number_format = '"$"#,##0'
            elif key == "sqft" and val is not None:
                cell.number_format = "#,##0"
            elif key in ("lat", "lon") and val is not None:
                cell.number_format = "0.0000"
            elif key == "available_from" and val is not None:
                cell.number_format = "yyyy-mm-dd"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)


if __name__ == "__main__":
    cli()
